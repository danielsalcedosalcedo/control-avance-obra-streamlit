"""
Módulo de Exportación a Excel con Formato Profesional
Genera un libro de Excel (.xlsx) estructurado en múltiples hojas con estilos aplicados.
"""

import io
from datetime import datetime
from typing import Dict, Any
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# Estilos de formato corporativo
COLOR_HEADER_BG = "1F4E79"       # Azul ejecutivo
COLOR_HEADER_FG = "FFFFFF"       # Blanco
COLOR_SUBHEADER_BG = "D9E1F2"    # Azul claro suave
COLOR_ZEBRA_BG = "F2F2F2"        # Gris claro alternado

FONT_TITLE = Font(name="Calibri", size=15, bold=True, color="1F4E79")
FONT_HEADER = Font(name="Calibri", size=11, bold=True, color=COLOR_HEADER_FG)
FONT_SUBHEADER = Font(name="Calibri", size=11, bold=True, color="000000")
FONT_BOLD = Font(name="Calibri", size=11, bold=True)
FONT_REGULAR = Font(name="Calibri", size=10)

FILL_HEADER = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
FILL_SUBHEADER = PatternFill(start_color=COLOR_SUBHEADER_BG, end_color=COLOR_SUBHEADER_BG, fill_type="solid")

ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

THIN_BORDER_SIDE = Side(border_style="thin", color="D3D3D3")
BORDER_THIN = Border(
    left=THIN_BORDER_SIDE,
    right=THIN_BORDER_SIDE,
    top=THIN_BORDER_SIDE,
    bottom=THIN_BORDER_SIDE
)


def _ajustar_ancho_columnas(hoja):
    """Ajusta automáticamente el ancho de cada columna según el contenido."""
    for col in hoja.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        hoja.column_dimensions[col_letter].width = max(max_len + 3, 12)


def _formatear_tabla_estandar(hoja, start_row: int, df: pd.DataFrame):
    """Aplica formato de cabecera, bordes y alineación a una tabla exportada."""
    num_cols = len(df.columns)
    
    # Formatear encabezados
    for col_idx in range(1, num_cols + 1):
        cell = hoja.cell(row=start_row, column=col_idx)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_THIN
    hoja.row_dimensions[start_row].height = 24

    # Formatear filas de datos
    for row_idx in range(start_row + 1, start_row + 1 + len(df)):
        hoja.row_dimensions[row_idx].height = 20
        for col_idx in range(1, num_cols + 1):
            cell = hoja.cell(row=row_idx, column=col_idx)
            cell.font = FONT_REGULAR
            cell.border = BORDER_THIN
            
            # Alineación según el tipo
            if isinstance(cell.value, (int, float)):
                cell.alignment = ALIGN_RIGHT
                if isinstance(cell.value, float):
                    cell.number_format = "0.00"
            else:
                cell.alignment = ALIGN_LEFT


def exportar_a_excel(
    df_base: pd.DataFrame,
    df_resumen_semanal: pd.DataFrame,
    df_especialidad: pd.DataFrame,
    df_encargados: pd.DataFrame,
    kpis: Dict[str, Any]
) -> bytes:
    """
    Genera un archivo Excel (.xlsx) en memoria con múltiples hojas estructuradas.
    
    Returns:
        Bytes del archivo Excel listo para st.download_button o guardado en disco.
    """
    output = io.BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        # 1. Hoja Resumen Ejecutivo
        df_kpi = pd.DataFrame([
            {"Indicador": "Fecha de Generación", "Valor": datetime.now().strftime("%Y-%m-%d %H:%M:%S")},
            {"Indicador": "Semanas Analizadas", "Valor": kpis.get("ultima_semana", 0)},
            {"Indicador": "Avance Planificado Acumulado", "Valor": f"{kpis.get('total_plan_acumulado', 0.0):.2f}%"},
            {"Indicador": "Avance Real Acumulado", "Valor": f"{kpis.get('total_real_acumulado', 0.0):.2f}%"},
            {"Indicador": "Desviación Global (Real - Plan)", "Valor": f"{kpis.get('desviacion_global', 0.0):.2f}%"},
            {"Indicador": "Estado de la Obra", "Valor": kpis.get("estado_global", "N/A")},
            {"Indicador": "SPI (Schedule Performance Index)", "Valor": f"{kpis.get('spi_global', 1.0):.2f}"},
            {"Indicador": "Total de Partidas Monitoreadas", "Valor": kpis.get("total_partidas", 0)},
            {"Indicador": "Total de Especialidades", "Valor": kpis.get("total_especialidades", 0)},
            {"Indicador": "Total de Responsables / Encargados", "Valor": kpis.get("total_encargados", 0)},
        ])
        df_kpi.to_excel(writer, sheet_name="Resumen_Ejecutivo", index=False, startrow=3)

        # 2. Hoja Control Semanal (Curva S)
        df_resumen_semanal.to_excel(writer, sheet_name="Control_Semanal", index=False)

        # 3. Hoja Especialidades
        df_especialidad.to_excel(writer, sheet_name="Por_Especialidad", index=False)

        # 4. Hoja Encargados
        df_encargados.to_excel(writer, sheet_name="Por_Encargado", index=False)

        # 5. Hoja Datos Detallados
        df_base.to_excel(writer, sheet_name="Datos_Detallados", index=False)

    # Reabrir con openpyxl para embellecer
    output.seek(0)
    workbook = openpyxl.load_workbook(output)

    # 1. Estilizar Resumen Ejecutivo
    hoja_resumen = workbook["Resumen_Ejecutivo"]
    hoja_resumen["A1"] = "INFORME EJECUTIVO DE CONTROL DE AVANCE FÍSICO DE OBRA"
    hoja_resumen["A1"].font = FONT_TITLE
    hoja_resumen["A2"] = "Métricas Clave de Rendimiento y Desempeño del Cronograma"
    hoja_resumen["A2"].font = FONT_SUBHEADER

    _formatear_tabla_estandar(hoja_resumen, start_row=4, df=df_kpi)
    _ajustar_ancho_columnas(hoja_resumen)

    # 2. Estilizar Control Semanal
    hoja_sem = workbook["Control_Semanal"]
    _formatear_tabla_estandar(hoja_sem, start_row=1, df=df_resumen_semanal)
    _ajustar_ancho_columnas(hoja_sem)

    # 3. Estilizar Por Especialidad
    hoja_esp = workbook["Por_Especialidad"]
    _formatear_tabla_estandar(hoja_esp, start_row=1, df=df_especialidad)
    _ajustar_ancho_columnas(hoja_esp)

    # 4. Estilizar Por Encargado
    hoja_enc = workbook["Por_Encargado"]
    _formatear_tabla_estandar(hoja_enc, start_row=1, df=df_encargados)
    _ajustar_ancho_columnas(hoja_enc)

    # 5. Estilizar Datos Detallados
    hoja_datos = workbook["Datos_Detallados"]
    _formatear_tabla_estandar(hoja_datos, start_row=1, df=df_base)
    _ajustar_ancho_columnas(hoja_datos)

    final_output = io.BytesIO()
    workbook.save(final_output)
    final_output.seek(0)
    return final_output.getvalue()
